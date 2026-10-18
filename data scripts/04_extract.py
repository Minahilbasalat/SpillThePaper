"""Step 4 - EXTRACT: read the text out of every downloaded file.

Reads : data/processed/collection_manifest.csv   (only rows with status "downloaded")
Saves : data/interim/text/<file_id>.txt           (extracted text, git-ignored)
        data/processed/extraction_results.csv     (one row per file: how extraction went)

Methods (chosen from what we saw in the real files):
  pdf_text    - normal PDF with real text            -> PyMuPDF
  docx        - Word file                            -> python-docx
  needs_ocr   - scanned PDF (almost no text) or a photo (.jpg/.png).
                We do NOT run OCR here; we only flag these files. OCR is a later,
                optional step once we know how many files need it.
  error       - file could not be opened (corrupt / password protected)

Usage:
    python "data scripts/04_extract.py"           # extract files not done yet
    python "data scripts/04_extract.py" --force   # redo everything
"""
import argparse

import pandas as pd

from common import PROCESSED_DIR, RAW_DIR, ROOT, log

try:
    import pymupdf  # new name of the library
except ImportError:  # older versions
    import fitz as pymupdf

import docx  # python-docx

TEXT_DIR = ROOT / "data" / "interim" / "text"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
# A text PDF normally has hundreds of characters per page. Below this we call it "scanned".
MIN_CHARS_PER_PAGE = 50


def extract_pdf(path):
    """Return (text, page_count) for a PDF."""
    with pymupdf.open(path) as document:
        pages = [page.get_text() for page in document]
    return "\n".join(pages), len(pages)


def extract_docx(path):
    """Return the text of a Word file (paragraphs + table cells)."""
    document = docx.Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_one(row):
    """Extract one file. Returns a dict with the result (and the text)."""
    path = RAW_DIR / row.local_file
    result = {"file_id": row.file_id, "extraction_method": "", "page_count": None,
              "char_count": 0, "word_count": 0, "needs_ocr": False, "error": "", "text": ""}
    try:
        if row.extension == ".pdf":
            text, pages = extract_pdf(path)
            result["page_count"] = pages
            if len(text.strip()) < MIN_CHARS_PER_PAGE * max(pages, 1):
                result["extraction_method"] = "needs_ocr"
                result["needs_ocr"] = True
            else:
                result["extraction_method"] = "pdf_text"
                result["text"] = text
        elif row.extension == ".docx":
            result["extraction_method"] = "docx"
            result["text"] = extract_docx(path)
        elif row.extension in IMAGE_EXTENSIONS:
            result["extraction_method"] = "needs_ocr"
            result["needs_ocr"] = True
            result["page_count"] = 1
        else:
            result["extraction_method"] = "error"
            result["error"] = f"unsupported extension {row.extension}"
    except Exception as error:  # corrupt or password-protected files
        result["extraction_method"] = "error"
        result["error"] = f"{type(error).__name__}: {str(error)[:100]}"
    result["char_count"] = len(result["text"].strip())
    result["word_count"] = len(result["text"].split())
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true", help="redo files already extracted")
    args = parser.parse_args()

    manifest = pd.read_csv(PROCESSED_DIR / "collection_manifest.csv", keep_default_na=False)
    downloaded = manifest[manifest["status"] == "downloaded"]
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    log(f"{len(downloaded)} downloaded files in the manifest", "extraction_log.txt")

    results_path = PROCESSED_DIR / "extraction_results.csv"
    old = (pd.read_csv(results_path, keep_default_na=False)
           if results_path.exists() and not args.force else pd.DataFrame())
    already_done = set(old["file_id"]) if len(old) else set()

    new_rows = []
    todo = [r for r in downloaded.itertuples() if r.file_id not in already_done]
    for number, row in enumerate(todo, start=1):
        result = extract_one(row)
        if result["text"]:
            (TEXT_DIR / f"{row.file_id}.txt").write_text(result["text"], encoding="utf-8")
        result.pop("text")
        new_rows.append(result)
        if number % 100 == 0 or number == len(todo):
            log(f"  progress: {number}/{len(todo)}", "extraction_log.txt")

    results = pd.concat([old, pd.DataFrame(new_rows)], ignore_index=True)
    results.to_csv(results_path, index=False, encoding="utf-8")

    log(f"Extraction methods: {results['extraction_method'].value_counts().to_dict()}",
        "extraction_log.txt")
    log(f"Files needing OCR: {int(results['needs_ocr'].sum())}", "extraction_log.txt")
    errors = results[results["error"] != ""]
    log(f"Files with errors: {len(errors)}", "extraction_log.txt")
    for row in errors.head(10).itertuples():
        log(f"  error {row.file_id}: {row.error}", "extraction_log.txt")
    log(f"Saved {results_path}", "extraction_log.txt")


if __name__ == "__main__":
    main()